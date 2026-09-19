import hashlib, io, os, re, tempfile
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import faiss, gdown, numpy as np, streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

APP_TITLE='AI Document Assistant'
EMBEDDING_MODEL_NAME='sentence-transformers/all-MiniLM-L6-v2'
GROQ_MODEL='openai/gpt-oss-120b'
CHUNK_SIZE=900
CHUNK_OVERLAP=150
SEMANTIC_TOP_K=8
KEYWORD_TOP_K=8
FINAL_TOP_K=6
SUPPORTED_EXTENSIONS={'.pdf','.docx','.txt','.md'}

st.set_page_config(page_title=APP_TITLE,page_icon='📚',layout='wide',initial_sidebar_state='expanded')
st.markdown('''<style>
.stApp{background:#f7f9fc}.block-container{max-width:1200px;padding-top:2rem;padding-bottom:3rem}
.hero{padding:1.5rem 1.7rem;border-radius:22px;background:linear-gradient(135deg,#111827,#243b53);color:white;margin-bottom:1.2rem;box-shadow:0 12px 35px rgba(15,23,42,.12)}
.hero h1{margin:0;font-size:2.1rem}.hero p{margin:.45rem 0 0;color:#dbeafe}.stat-card{background:white;border:1px solid #e5e7eb;border-radius:16px;padding:1rem 1.1rem;box-shadow:0 5px 18px rgba(15,23,42,.05)}
.stat-label{color:#64748b;font-size:.82rem}.stat-value{color:#0f172a;font-size:1.45rem;font-weight:700}.source-card{background:white;border:1px solid #e2e8f0;border-left:4px solid #2563eb;border-radius:12px;padding:.85rem 1rem;margin:.55rem 0}.source-meta{color:#64748b;font-size:.78rem;margin-bottom:.35rem}.source-text{color:#334155;font-size:.9rem;line-height:1.55}div[data-testid="stFileUploader"]{background:white;border-radius:16px;padding:.35rem}
</style>''',unsafe_allow_html=True)

def init_state():
    defaults={'chunks':[],'embeddings':None,'faiss_index':None,'documents':{},'processed_hashes':set(),'messages':[],'drive_files_loaded':[]}
    for k,v in defaults.items():
        if k not in st.session_state: st.session_state[k]=v
init_state()

@st.cache_resource(show_spinner='Loading embedding model...')
def get_embedding_model(): return SentenceTransformer(EMBEDDING_MODEL_NAME)
@st.cache_resource
def get_groq_client(api_key): return Groq(api_key=api_key)

def clean_text(text):
    text=text.replace('\x00',' ')
    return '\n'.join(re.sub(r'[ \t]+',' ',x).strip() for x in text.splitlines() if x.strip())
def decode_text(data):
    for enc in ('utf-8','utf-8-sig','cp1252','latin-1'):
        try:return data.decode(enc)
        except UnicodeDecodeError:pass
    return data.decode('utf-8',errors='ignore')

def extract_pdf(data,filename):
    reader=PdfReader(io.BytesIO(data)); out=[]
    for n,page in enumerate(reader.pages,1):
        text=clean_text(page.extract_text() or '')
        if text: out.append({'text':text,'filename':filename,'page':n})
    return out

def extract_docx(data,filename):
    doc=Document(io.BytesIO(data)); parts=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells=[c.text.strip() for c in row.cells if c.text.strip()]
            if cells: parts.append(' | '.join(cells))
    text=clean_text('\n'.join(parts)); return [{'text':text,'filename':filename,'page':None}] if text else []
def extract_txt(data,filename):
    text=clean_text(decode_text(data)); return [{'text':text,'filename':filename,'page':None}] if text else []
def extract_md(data,filename): return extract_txt(data,filename)
def extract_document(data,filename):
    ext=Path(filename).suffix.lower()
    return {'.pdf':extract_pdf,'.docx':extract_docx,'.txt':extract_txt,'.md':extract_md}.get(ext,lambda *_:[])(data,filename)

def chunk_documents(pages):
    chunks=[]; cid=0
    for page in pages:
        text=page['text'].strip(); start=0
        while start<len(text):
            end=min(start+CHUNK_SIZE,len(text)); piece=text[start:end].strip()
            if piece:
                chunks.append({'id':cid,'text':piece,'filename':page['filename'],'page':page.get('page')}); cid+=1
            if end>=len(text): break
            start=max(end-CHUNK_OVERLAP,start+1)
    return chunks

def add_chunks_to_index(new_chunks):
    if not new_chunks:return 0
    model=get_embedding_model(); texts=[x['text'] for x in new_chunks]
    emb=model.encode(texts,batch_size=32,show_progress_bar=False,convert_to_numpy=True,normalize_embeddings=True).astype('float32')
    if st.session_state['faiss_index'] is None: st.session_state['faiss_index']=faiss.IndexFlatIP(emb.shape[1])
    st.session_state['faiss_index'].add(emb)
    st.session_state['embeddings']=emb if st.session_state['embeddings'] is None else np.vstack([st.session_state['embeddings'],emb])
    st.session_state['chunks'].extend(new_chunks); return len(new_chunks)

def file_hash(data): return hashlib.sha256(data).hexdigest()

def process_uploaded_files(uploaded_files):
    docs=chunks=skipped=0; errors=[]
    for f in uploaded_files or []:
        data=f.getvalue(); digest=file_hash(data)
        if digest in st.session_state['processed_hashes']: skipped+=1; continue
        try:
            pages=extract_document(data,f.name); new=chunk_documents(pages)
            if not new: errors.append(f'{f.name}: no extractable text'); continue
            offset=len(st.session_state['chunks'])
            for i,x in enumerate(new): x['id']=offset+i
            added=add_chunks_to_index(new); st.session_state['processed_hashes'].add(digest)
            st.session_state['documents'][f.name]={'filename':f.name,'pages':len(pages),'chunks':added}; docs+=1; chunks+=added
        except Exception as e: errors.append(f'{f.name}: {e}')
    return docs,chunks,skipped,errors

def is_drive_folder_url(url): return '/folders/' in urlparse(url).path

def download_drive_source(url):
    temp=Path(tempfile.mkdtemp(prefix='drive_docs_'))
    if is_drive_folder_url(url):
        result=gdown.download_folder(url,output=str(temp),quiet=True,use_cookies=False)
        return [Path(x) for x in (result or []) if Path(x).is_file()]
    result=gdown.download(url=url,output=str(temp/'drive_download'),quiet=True,use_cookies=False)
    return [Path(result)] if result and Path(result).is_file() else []

def process_drive_url(url):
    if not url.strip(): return 0,0,[], 'Please paste a Google Drive file or folder link.'
    if 'drive.google.com' not in url: return 0,0,[],'Please provide a Google Drive link.'
    try: paths=download_drive_source(url)
    except Exception as e: return 0,0,[],f'Google Drive download failed: {e}'
    supported=[p for p in paths if p.suffix.lower() in SUPPORTED_EXTENSIONS]
    if not supported: return 0,0,[],'No supported PDF, DOCX, TXT or MD files were found. Make sure the shared file/folder contains supported files and is accessible.'
    docs=chunks=0; names=[]; errors=[]
    for p in supported:
        try:
            data=p.read_bytes(); digest=file_hash(data)
            if digest in st.session_state['processed_hashes']: continue
            pages=extract_document(data,p.name); new=chunk_documents(pages)
            if not new: errors.append(f'{p.name}: no extractable text'); continue
            offset=len(st.session_state['chunks'])
            for i,x in enumerate(new): x['id']=offset+i
            added=add_chunks_to_index(new); st.session_state['processed_hashes'].add(digest)
            st.session_state['documents'][p.name]={'filename':p.name,'pages':len(pages),'chunks':added}; docs+=1; chunks+=added; names.append(p.name)
        except Exception as e: errors.append(f'{p.name}: {e}')
    return docs,chunks,names,'\n'.join(errors)

def important_query_words(question):
    words=re.findall(r'\b[a-zA-Z0-9][a-zA-Z0-9+#.-]*\b',question.lower())
    stop={'the','a','an','is','are','was','were','what','which','who','when','where','why','how','does','do','did','can','could','would','should','and','or','of','to','in','on','for','with','from','about','this','that','these','those','it','its','be','as','by','at','has','have','had','i','me','my','we','you','your'}
    return [w for w in words if w not in stop and len(w)>2]

def semantic_search(question,top_k=SEMANTIC_TOP_K):
    index=st.session_state['faiss_index']; chunks=st.session_state['chunks']
    if index is None or not chunks:return {}
    vec=get_embedding_model().encode([question],convert_to_numpy=True,normalize_embeddings=True).astype('float32')
    scores,idx=index.search(vec,min(top_k,len(chunks))); return {int(i):float(s) for s,i in zip(scores[0],idx[0]) if i>=0}

def keyword_search(question,top_k=KEYWORD_TOP_K):
    chunks=st.session_state['chunks']; q=set(important_query_words(question)); scored=[]
    if not q:return {}
    for i,c in enumerate(chunks):
        words=set(re.findall(r'\b[a-zA-Z0-9][a-zA-Z0-9+#.-]*\b',c['text'].lower())); matches=len(q & words)
        if matches: scored.append((i,matches/len(q)))
    scored.sort(key=lambda x:x[1],reverse=True); return dict(scored[:top_k])

def hybrid_search(question,top_k=FINAL_TOP_K):
    semantic=semantic_search(question); keyword=keyword_search(question); ids=set(semantic)|set(keyword); ranked=[]
    for i in ids:
        ss=semantic.get(i,0); ks=keyword.get(i,0); ranked.append({'index':i,'score':.70*ss+.30*ks,'semantic_score':ss,'keyword_score':ks,'chunk':st.session_state['chunks'][i]})
    ranked.sort(key=lambda x:x['score'],reverse=True); return ranked[:top_k]

def get_api_key():
    try:return st.secrets['GROQ_API_KEY']
    except Exception:return None

def answer_question(question,results):
    if not results:return "I don't know based on the loaded documents."
    key=get_api_key()
    if not key: raise RuntimeError('GROQ_API_KEY is missing. Add it to Streamlit Secrets.')
    context='\n\n'.join(f"[Source {n}: {r['chunk']['filename']}" + (f" — Page {r['chunk']['page']}]" if r['chunk']['page'] else ']') + f"\n{r['chunk']['text']}" for n,r in enumerate(results,1))
    prompt=f'''Answer the question ONLY from the supplied document context. Do not use outside knowledge or invent facts. If the answer is unavailable, say exactly: "I don't know based on the loaded documents." Be clear and concise.\n\nDOCUMENT CONTEXT:\n{context}\n\nQUESTION:\n{question}'''
    try:
        response=get_groq_client(key).chat.completions.create(model=GROQ_MODEL,messages=[{'role':'system','content':'Answer only from supplied document context. Never fabricate missing information.'},{'role':'user','content':prompt}],temperature=0)
    except Exception as e:
        msg=str(e)
        if 'not found' in msg.lower() and 'model' in msg.lower(): raise RuntimeError(f"Groq model '{GROQ_MODEL}' was not found. Update GROQ_MODEL in app.py to a currently supported Groq model.") from e
        raise RuntimeError(f'Groq request failed: {e}') from e
    return response.choices[0].message.content

def render_sources(results):
    st.markdown('#### 📚 Retrieved Sources')
    for r in results:
        c=r['chunk']; page=f"Page {c['page']}" if c['page'] else 'Page not available'
        st.markdown(f'''<div class="source-card"><div class="source-meta"><strong>{c['filename']}</strong> • {page} • Hybrid score: {r['score']:.3f}</div><div class="source-text">{c['text']}</div></div>''',unsafe_allow_html=True)

st.markdown('<div class="hero"><h1>📚 AI Document Assistant</h1><p>Search your own PDF, DOCX, TXT and Markdown files with semantic + keyword retrieval, then ask grounded questions.</p></div>',unsafe_allow_html=True)
with st.sidebar:
    st.header('📥 Add Documents')
    files=st.file_uploader('Upload PDF, DOCX, TXT or MD',type=['pdf','docx','txt','md'],accept_multiple_files=True)
    if st.button('Process uploaded files',use_container_width=True):
        with st.spinner('Extracting, chunking and embedding...'): d,c,s,e=process_uploaded_files(files)
        if d: st.success(f'Added {d} document(s) and {c} chunk(s).')
        elif s: st.info('Selected files were already processed.')
        for x in e: st.warning(x)
    st.divider(); st.subheader('☁️ Google Drive')
    drive=st.text_input('Drive file or folder link',placeholder='https://drive.google.com/...')
    if st.button('Load from Google Drive',use_container_width=True):
        with st.spinner('Downloading and processing Drive files...'): d,c,n,msg=process_drive_url(drive)
        if d: st.success(f'Loaded {d} file(s) and created {c} chunk(s).')
        if msg: (st.caption(msg) if d else st.error(msg))
    st.divider(); st.subheader('⚙️ Session')
    st.write(f"Documents: **{len(st.session_state['documents'])}**"); st.write(f"Chunks: **{len(st.session_state['chunks'])}**")
    if st.button('🗑️ Clear documents',use_container_width=True):
        for k,v in {'chunks':[],'embeddings':None,'faiss_index':None,'documents':{},'processed_hashes':set(),'messages':[],'drive_files_loaded':[]}.items(): st.session_state[k]=v
        st.rerun()

if st.session_state['documents']:
    cols=st.columns(3)
    for col,(label,val) in zip(cols,[('📄 Documents',len(st.session_state['documents'])),('🧩 Chunks',len(st.session_state['chunks'])),('🔎 Retrieval','Hybrid')]):
        col.markdown(f'<div class="stat-card"><div class="stat-label">{label}</div><div class="stat-value">{val}</div></div>',unsafe_allow_html=True)
    with st.expander('📑 Loaded documents'):
        for name,info in st.session_state['documents'].items(): st.write(f"**{name}** — {info['pages']} pages — {info['chunks']} chunks")

st.divider(); st.subheader('💬 Ask your documents')
if not st.session_state['chunks']: st.info('Upload a document or load a supported Google Drive file/folder to start.')
for m in st.session_state['messages']:
    with st.chat_message(m['role']):
        st.markdown(m['content'])
        if m.get('sources'): render_sources(m['sources'])

question=st.chat_input('Ask a question about your loaded documents...')
if question:
    if not st.session_state['chunks']: st.warning('Please load at least one document first.'); st.stop()
    st.session_state['messages'].append({'role':'user','content':question})
    with st.chat_message('user'): st.markdown(question)
    with st.chat_message('assistant'):
        try:
            with st.spinner('Searching your documents...'): results=hybrid_search(question)
            with st.spinner('Generating answer...'): answer=answer_question(question,results)
            st.markdown(answer)
            if results: render_sources(results)
            st.session_state['messages'].append({'role':'assistant','content':answer,'sources':results})
        except Exception as e:
            msg=f'⚠️ {e}'; st.error(msg); st.session_state['messages'].append({'role':'assistant','content':msg,'sources':results if 'results' in locals() else []})
